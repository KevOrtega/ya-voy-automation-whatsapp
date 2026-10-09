"""``.xlsx`` builder (openpyxl, in memory; docs/05 "Contenido del archivo", docs/06 ``reports``).

Bold headers, date columns as real Excel datetimes in ``TZ`` local time with the
``dd/mm/yyyy hh:mm`` format, auto filters. Corrections add the ``Correcciones`` sheet, the
``Cambio`` column in ``Pedidos`` and the "CORRECCIÓN — versión N" banner on the first row of
every sheet. Pure: no database, no network.
"""

from __future__ import annotations

import datetime as dt
from collections.abc import Sequence
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ya_voy.messaging import catalog
from ya_voy.messaging import report_texts as texts
from ya_voy.reports.model import ReportData

DATETIME_FORMAT = "dd/mm/yyyy hh:mm"
BOLD = Font(bold=True)
SUMMARY_WIDTHS = (24, 16, 10, 10, 11)
ORDERS_WIDTHS = (10, 17, 17, 22, 30, 30, 24, 16, 20, 11, 18)
CORRECTIONS_WIDTHS = (10, 22, 22, 60)


def local_naive(at: dt.datetime | None, tz: ZoneInfo) -> dt.datetime | None:
    """Excel has no time zones: the cell holds the wall-clock time in ``tz``."""
    if at is None:
        return None
    return at.astimezone(tz).replace(tzinfo=None)


def _table(
    ws: Worksheet,
    data: ReportData,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    widths: Sequence[int],
    datetime_columns: Sequence[int] = (),
) -> None:
    if data.is_correction:
        ws.append([catalog.CORRECTION_BANNER.format(revision=data.revision)])
        ws.cell(row=1, column=1).font = BOLD
    header_row = ws.max_row + 1 if data.is_correction else 1
    ws.append(list(headers))
    for cell in ws[header_row]:
        cell.font = BOLD
    for row in rows:
        ws.append(list(row))
    for column in datetime_columns:
        for (cell,) in ws.iter_rows(
            min_row=header_row + 1, min_col=column, max_col=column, max_row=ws.max_row
        ):
            cell.number_format = DATETIME_FORMAT
    last = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{header_row}:{last}{max(ws.max_row, header_row)}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for index, width in enumerate(widths[: len(headers)], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def build_workbook(data: ReportData, tz: ZoneInfo) -> bytes:
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    summary.title = catalog.SHEET_SUMMARY
    _table(
        summary,
        data,
        catalog.SUMMARY_HEADERS,
        [
            (c.name, c.phone, c.taken, c.released, c.cancelled)
            for c in sorted(data.couriers, key=lambda c: (c.name.casefold(), c.courier_id))
        ],
        SUMMARY_WIDTHS,
    )

    changes = data.change_by_code()
    headers = list(catalog.ORDERS_HEADERS)
    if data.is_correction:
        headers.append(catalog.ORDERS_CHANGE_HEADER)
    order_rows = []
    for o in sorted(data.orders, key=lambda o: (o.created_at, o.code)):
        assigned = o.status == "assigned"
        row: list[Any] = [
            o.code,
            local_naive(o.created_at, tz),
            local_naive(o.assigned_at, tz) if assigned else None,
            o.store_name,
            o.pickup_address,
            o.delivery_address,
            o.product,
            o.payment_label,
            o.courier_name if assigned else None,
            catalog.ORDER_STATUS_LABELS[o.status],
        ]
        if data.is_correction:
            change = changes.get(o.code)
            row.append(catalog.CHANGE_LABELS[change.type] if change else None)
        order_rows.append(row)
    _table(
        wb.create_sheet(catalog.SHEET_ORDERS),
        data,
        headers,
        order_rows,
        ORDERS_WIDTHS,
        datetime_columns=(2, 3),
    )

    if data.is_correction:
        _table(
            wb.create_sheet(catalog.SHEET_CORRECTIONS),
            data,
            catalog.CORRECTIONS_HEADERS,
            [
                (c.code, texts.before_label(c), texts.after_label(c), texts.change_detail(c))
                for c in data.changes
            ],
            CORRECTIONS_WIDTHS,
        )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
