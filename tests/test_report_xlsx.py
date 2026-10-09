"""The generated workbook, reopened with openpyxl (docs/05, RF-51, RF-52)."""

from __future__ import annotations

import datetime as dt
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from tests.report_helpers import courier_row, order_row
from ya_voy.reports.model import Change, ChangeType, ReportData
from ya_voy.reports.xlsx import build_workbook

BA = ZoneInfo("America/Argentina/Buenos_Aires")
UTC = dt.UTC
D = dt.date(2026, 10, 9)


def utc(day: int, hour: int, minute: int = 0) -> dt.datetime:
    return dt.datetime(2026, 10, day, hour, minute, tzinfo=UTC)


def reopen(data: ReportData):
    return load_workbook(BytesIO(build_workbook(data, BA)))


def values(ws) -> list[tuple]:
    return [tuple(row) for row in ws.iter_rows(values_only=True)]


JOEL = (7, "Joel")
EMI = (8, "Emiliano")


def initial() -> ReportData:
    return ReportData(
        report_date=D,
        revision=1,
        orders=[
            order_row(
                2,
                "YV-0124",
                "cancelled",
                created_at=utc(9, 22, 30),  # 19:30 local
                store_name="Almacén Don Luis",
                payment="efectivo",
            ),
            order_row(
                1,
                "YV-0123",
                "assigned",
                courier=JOEL,
                created_at=utc(9, 15, 0),  # 12:00 local
                assigned_at=utc(9, 15, 4),
            ),
        ],
        couriers=[
            courier_row(8, "Emiliano", released=1, cancelled=1),
            courier_row(7, "Joel", phone="5491100000007", taken=1),
        ],
    )


def test_initial_report_sheets_headers_and_rows() -> None:
    wb = reopen(initial())
    assert wb.sheetnames == ["Resumen", "Pedidos"]

    assert values(wb["Resumen"]) == [
        ("Motero", "Teléfono", "Tomados", "Liberados", "Cancelados"),
        ("Emiliano", None, 0, 1, 1),  # known only by @lid: no phone
        ("Joel", "5491100000007", 1, 0, 0),
    ]
    rows = values(wb["Pedidos"])
    assert rows[0] == (
        "Código",
        "Creado",
        "Asignado",
        "Local",
        "Retiro",
        "Entrega",
        "Producto",
        "Medio de pago",
        "Motero",
        "Estado",
    )
    assert rows[1:] == [
        (
            "YV-0123",
            dt.datetime(2026, 10, 9, 12, 0),
            dt.datetime(2026, 10, 9, 12, 4),
            "Casa Nogal",
            "casa nogal almendros",
            "el ceibo 2871",
            "almacen",
            "transferencia",
            "Joel",  # the courier's name, never an id or JID
            "Tomado",
        ),
        (
            "YV-0124",
            dt.datetime(2026, 10, 9, 19, 30),
            None,
            "Almacén Don Luis",
            "casa nogal almendros",
            "el ceibo 2871",
            "almacen",
            "efectivo",
            None,
            "Cancelado",
        ),
    ]


def test_format_bold_headers_filters_and_datetime_format() -> None:
    wb = reopen(initial())
    orders = wb["Pedidos"]
    assert all(cell.font.bold for cell in orders[1])
    assert orders.auto_filter.ref == "A1:J3"
    assert wb["Resumen"].auto_filter.ref == "A1:E3"
    assert orders["B2"].number_format == "dd/mm/yyyy hh:mm"
    assert orders["C2"].number_format == "dd/mm/yyyy hh:mm"
    assert orders["B2"].is_date


def test_utc_times_around_the_0500_rollover_show_in_buenos_aires_time() -> None:
    """Stored UTC 07:59 and 08:01 of 10/10 are 04:59 and 05:01 in Buenos Aires (UTC-3). Both
    rows render in local time; which report they belong to is ``business_date``'s job."""
    data = ReportData(
        report_date=D,
        revision=1,
        orders=[
            order_row(
                1,
                "YV-0201",
                "assigned",
                courier=JOEL,
                created_at=utc(10, 2, 55),  # 23:55 of 09/10 local
                assigned_at=utc(10, 7, 59),  # 04:59 of 10/10 local
            ),
            order_row(
                2,
                "YV-0202",
                "assigned",
                courier=EMI,
                created_at=utc(10, 7, 58),
                assigned_at=utc(10, 8, 1),  # 05:01 local
            ),
        ],
        couriers=[courier_row(7, "Joel", taken=1), courier_row(8, "Emiliano", taken=1)],
    )
    rows = values(reopen(data)["Pedidos"])[1:]
    assert [(r[0], r[1], r[2]) for r in rows] == [
        ("YV-0201", dt.datetime(2026, 10, 9, 23, 55), dt.datetime(2026, 10, 10, 4, 59)),
        ("YV-0202", dt.datetime(2026, 10, 10, 4, 58), dt.datetime(2026, 10, 10, 5, 1)),
    ]


def test_empty_day_has_sheets_with_headers_only() -> None:
    wb = reopen(ReportData(report_date=D, revision=1, orders=[], couriers=[]))
    assert values(wb["Resumen"]) == [("Motero", "Teléfono", "Tomados", "Liberados", "Cancelados")]
    assert len(values(wb["Pedidos"])) == 1
    assert wb["Pedidos"].auto_filter.ref == "A1:J1"


def test_correction_has_banner_change_column_and_corrections_sheet() -> None:
    data = ReportData(
        report_date=D,
        revision=2,
        orders=[
            order_row(1, "YV-0123", "assigned", courier=EMI, created_at=utc(9, 15)),
            order_row(2, "YV-0124", "cancelled", created_at=utc(9, 16)),
            order_row(3, "YV-0130", "assigned", courier=JOEL, created_at=utc(10, 2, 40)),
            order_row(4, "YV-0131", "cancelled", created_at=utc(10, 2, 45)),
            order_row(5, "YV-0125", "assigned", courier=JOEL, created_at=utc(9, 17)),
        ],
        couriers=[courier_row(8, "Emiliano", taken=1), courier_row(7, "Joel", taken=2)],
        changes=[
            Change("YV-0123", ChangeType.COURIER_CHANGED, "Joel", "Emiliano"),
            Change("YV-0124", ChangeType.CANCELLED, "Joel", None),
            Change("YV-0130", ChangeType.ADDED, None, "Joel"),
            Change("YV-0131", ChangeType.ADDED, None, None),
        ],
    )
    wb = reopen(data)
    assert wb.sheetnames == ["Resumen", "Pedidos", "Correcciones"]
    for name in wb.sheetnames:
        assert wb[name]["A1"].value == "CORRECCIÓN — versión 2"
        assert wb[name]["A1"].font.bold
        assert all(cell.font.bold for cell in wb[name][2])

    orders = values(wb["Pedidos"])
    assert orders[1][-1] == "Cambio"
    assert [(r[0], r[8], r[9], r[10]) for r in orders[2:]] == [
        ("YV-0123", "Emiliano", "Tomado", "Cambió de motero"),
        ("YV-0124", None, "Cancelado", "Cancelado"),
        ("YV-0125", "Joel", "Tomado", None),
        ("YV-0130", "Joel", "Tomado", "Agregado"),
        ("YV-0131", None, "Cancelado", "Agregado"),
    ]
    assert wb["Pedidos"].auto_filter.ref == "A2:K7"
    assert values(wb["Correcciones"])[1:] == [
        ("Código", "Antes", "Ahora", "Detalle"),
        (
            "YV-0123",
            "Tomado por Joel",
            "Tomado por Emiliano",
            "El pedido YV-0123 pasó de Joel a Emiliano.",
        ),
        (
            "YV-0124",
            "Tomado por Joel",
            "Cancelado",
            "El pedido YV-0124 fue cancelado (estaba asignado a Joel).",
        ),
        ("YV-0130", "No estaba", "Tomado por Joel", "Pedido agregado YV-0130 (tomado por Joel)"),
        ("YV-0131", "No estaba", "Cancelado", "Pedido agregado YV-0131 (cancelado)"),
    ]
    assert values(wb["Resumen"])[2:] == [("Emiliano", None, 1, 0, 0), ("Joel", None, 2, 0, 0)]
