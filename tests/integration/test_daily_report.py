"""Daily report and corrections against Postgres/Redis (CU-16, CU-17, RF-48..RF-55).

"Now" is frozen through ``schedule.repository.db_now``; Evolution is mocked with respx. The helpers
below apply the documented SQL (docs/03) plus ``mark_report_for_correction`` in one transaction, to
control lock timing precisely; the real transitions (``ya_voy.orders``, stage 5a/6a) do the same and
their report effect is covered in ``test_orders.py``.
"""

from __future__ import annotations

import asyncio
import datetime as dt
import json
import logging
from collections.abc import Callable, Iterator
from io import BytesIO
from typing import Any

import httpx
import pytest
import respx
from openpyxl import load_workbook
from sqlalchemy import select, text, update
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession, async_sessionmaker

from scripts import resend_report
from tests.integration import factories as f
from tests.integration.conftest import EVOLUTION_URL, queued, redis_test_url
from tests.schedule_helpers import at
from ya_voy.db import models as m
from ya_voy.queue import SEND_OUTBOUND
from ya_voy.reports import service
from ya_voy.reports.service import Outcome
from ya_voy.schedule import repository as schedule_repository
from ya_voy.worker import reports as report_jobs
from ya_voy.worker.outbound import send_outbound

Ctx = Callable[..., dict[str, Any]]
SF = async_sessionmaker[AsyncSession]
D = dt.date(2026, 10, 9)
D1 = D + dt.timedelta(days=1)
ADMINS = ("5491100000001", "5491100000002")


class Clock:
    def __init__(self, now: dt.datetime) -> None:
        self.now = now

    async def __call__(self, session: AsyncSession) -> dt.datetime:
        return self.now


@pytest.fixture
def clock(monkeypatch: pytest.MonkeyPatch) -> Clock:
    c = Clock(at(9, 22))
    monkeypatch.setattr(schedule_repository, "db_now", c)
    return c


@pytest.fixture
def evo() -> Iterator[respx.MockRouter]:
    with respx.mock(base_url=EVOLUTION_URL, assert_all_called=False) as router:
        counter = iter(range(1, 10_000))
        for path, name in (
            ("/message/sendMedia/yavoy", "sendMedia"),
            ("/message/sendText/yavoy", "sendText"),
        ):
            router.post(path, name=name).mock(
                side_effect=lambda request: httpx.Response(
                    201, json={"key": {"id": f"BAE5{next(counter)}", "fromMe": True}}
                )
            )
        yield router


@pytest.fixture
def ctx(worker_ctx: Ctx, clock: Clock, evo: respx.MockRouter) -> dict[str, Any]:
    return worker_ctx()


# --- data helpers (``ya_voy.orders`` writes these rows for real) -------------------------------


class World:
    def __init__(self, engine: AsyncEngine) -> None:
        self.engine = engine
        self.store_id = 0
        self.couriers: dict[str, int] = {}

    async def setup(self) -> World:
        self.store_id = await f.insert_ok(self.engine, f.store(display_name="Casa Nogal"))
        self.couriers["Joel"] = await f.insert_ok(
            self.engine,
            f.courier(
                display_name="Joel", phone="5491100000007", phone_jid="5491100000007@s.whatsapp.net"
            ),
        )
        self.couriers["Emiliano"] = await f.insert_ok(
            self.engine,
            f.courier(display_name="Emiliano"),  # only known by @lid
        )
        return self

    async def order(
        self,
        code: str | None,
        status: str,
        *,
        day: dt.date = D,
        created: dt.datetime | None = None,
        courier: str | None = None,
        assigned: dt.datetime | None = None,
    ) -> int:
        created = created or at(9, 15)
        courier_id = self.couriers[courier] if courier else None
        values: dict[str, Any] = {
            "status": status,
            "business_date": day,
            "created_at": created,
            "pickup_address": "casa nogal almendros",
            "delivery_address": "el ceibo 2871",
            "product": "almacen",
            "payment_method": "transfer",
        }
        if code:
            values |= {"code": code, "published_at": created + dt.timedelta(minutes=1)}
        if status == "assigned":
            values |= {
                "courier_id": courier_id,
                "last_courier_id": courier_id,
                "assigned_at": assigned or created + dt.timedelta(minutes=3),
            }
        order_id = await f.insert_ok(self.engine, f.order(self.store_id, **values))
        if status == "assigned":
            await self.event(order_id, "claimed", courier, "published", "assigned")
        return order_id

    async def event(
        self, order_id: int, kind: str, courier: str | None, from_status: str, to_status: str
    ) -> None:
        await f.insert_ok(
            self.engine,
            f.insert(m.OrderEvent).values(
                order_id=order_id,
                event_type=kind,
                actor_type="courier",
                courier_id=self.couriers[courier] if courier else None,
                from_status=from_status,
                to_status=to_status,
            ),
        )

    async def release(self, session: AsyncSession, order_id: int, courier: str) -> None:
        """T13 + docs/03 "Efecto sobre el reporte diario", in the caller's transaction."""
        await session.execute(
            update(m.Order)
            .where(m.Order.id == order_id, m.Order.status == "assigned")
            .values(courier_id=None, status="published", assigned_at=None)
        )
        session.add(
            m.OrderEvent(
                order_id=order_id,
                event_type="released_by_courier",
                actor_type="courier",
                courier_id=self.couriers[courier],
                from_status="assigned",
                to_status="published",
            )
        )
        await service.mark_report_for_correction(session, D)

    async def claim(self, order_id: int, courier: str, when: dt.datetime) -> None:
        async with self.engine.begin() as conn:
            await conn.execute(
                update(m.Order)
                .where(m.Order.id == order_id, m.Order.status == "published")
                .values(
                    courier_id=self.couriers[courier],
                    last_courier_id=self.couriers[courier],
                    status="assigned",
                    assigned_at=when,
                )
            )
        await self.event(order_id, "claimed", courier, "published", "assigned")

    async def cancel(self, session: AsyncSession, order_id: int, courier: str) -> None:
        await session.execute(
            update(m.Order)
            .where(m.Order.id == order_id)
            .values(
                status="cancelled",
                courier_id=None,
                cancelled_at=at(9, 23, 30),
                cancelled_by="admin",
            )
        )
        session.add(
            m.OrderEvent(
                order_id=order_id,
                event_type="cancelled",
                actor_type="admin",
                courier_id=self.couriers[courier],
                from_status="assigned",
                to_status="cancelled",
            )
        )
        await service.mark_report_for_correction(session, D)


@pytest.fixture
async def world(engine: AsyncEngine) -> World:
    return await World(engine).setup()


async def opening(
    engine: AsyncEngine, start: dt.datetime, end: dt.datetime, confirmed: dt.datetime
) -> int:
    return await f.insert_ok(
        engine,
        f.override(
            kind="open_window",
            status="active",
            business_date=D,
            starts_at=start,
            ends_at=end,
            confirmed_at=confirmed,
            source_text="hoy hasta la 01",
            parsed={"kind": "open_window"},
        ),
    )


async def report(sf: SF, day: dt.date = D) -> m.DailyReport | None:
    async with sf() as session:
        return await session.scalar(select(m.DailyReport).where(m.DailyReport.report_date == day))


async def revisions(sf: SF) -> list[m.ReportRevision]:
    async with sf() as session:
        return list(
            await session.scalars(select(m.ReportRevision).order_by(m.ReportRevision.revision))
        )


async def outbox(sf: SF, template: str | None = None) -> list[m.OutboundMessage]:
    async with sf() as session:
        rows = await session.scalars(select(m.OutboundMessage).order_by(m.OutboundMessage.id))
        return [r for r in rows if template is None or r.template_key == template]


async def deliver(ctx: dict[str, Any]) -> list[str]:
    results = []
    for row in await outbox(ctx["session_factory"]):
        if row.status == "pending":
            results.append(await send_outbound(ctx, row.id))
    return results


def workbook(rev: m.ReportRevision):
    assert rev.file_content is not None
    return load_workbook(BytesIO(rev.file_content))


def rows(ws) -> list[tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


async def run(ctx: dict[str, Any], day: dt.date = D) -> Outcome:
    outcome, _ = await service.run_step(ctx["session_factory"], ctx["settings"], day)
    return outcome


# --- CU-16: the original report -----------------------------------------------------------------


async def test_report_at_effective_close_with_names_and_local_times(
    ctx: dict[str, Any], clock: Clock, world: World, evo: respx.MockRouter
) -> None:
    sf = ctx["session_factory"]
    o1 = await world.order(
        "YV-0123", "assigned", courier="Joel", created=at(9, 12), assigned=at(9, 12, 4)
    )
    await world.order("YV-0124", "assigned", courier="Emiliano", created=at(9, 20, 30))
    o3 = await world.order("YV-0125", "cancelled", created=at(9, 21))
    await world.event(o3, "claimed", "Joel", "published", "assigned")
    await world.event(o3, "cancelled", "Joel", "assigned", "cancelled")
    await world.order(None, "expired", created=at(9, 21, 30))  # never published: not reported
    await world.event(o1, "released_by_courier", "Emiliano", "assigned", "published")

    clock.now = at(9, 22, 59)
    assert await run(ctx) == Outcome.NOT_DUE
    assert await report(sf) is None

    clock.now = at(9, 23)
    assert await run(ctx) == Outcome.STARTED
    [rev] = await revisions(sf)
    assert (rev.revision, rev.kind, rev.status, rev.attempts) == (1, "initial", "generated", 1)
    assert rev.file_name == "reporte-ya-voy-2026-10-09.xlsx"
    assert rev.snapshot == {
        str(o1): {
            "code": "YV-0123",
            "status": "assigned",
            "courier_id": world.couriers["Joel"],
            "courier_name": "Joel",
        },
        str(o1 + 1): {
            "code": "YV-0124",
            "status": "assigned",
            "courier_id": world.couriers["Emiliano"],
            "courier_name": "Emiliano",
        },
        str(o3): {
            "code": "YV-0125",
            "status": "cancelled",
            "courier_id": None,
            "courier_name": None,
        },
    }

    wb = workbook(rev)
    assert rows(wb["Resumen"]) == [
        ("Motero", "Teléfono", "Tomados", "Liberados", "Cancelados"),
        ("Emiliano", None, 1, 1, 0),
        ("Joel", "5491100000007", 1, 0, 1),
    ]
    orders = rows(wb["Pedidos"])
    assert len(orders) == 4
    assert [(r[0], r[1], r[2], r[3], r[7], r[8], r[9]) for r in orders[1:]] == [
        (
            "YV-0123",
            dt.datetime(2026, 10, 9, 12, 0),
            dt.datetime(2026, 10, 9, 12, 4),
            "Casa Nogal",
            "transferencia",
            "Joel",
            "Tomado",
        ),
        (
            "YV-0124",
            dt.datetime(2026, 10, 9, 20, 30),
            dt.datetime(2026, 10, 9, 20, 33),
            "Casa Nogal",
            "transferencia",
            "Emiliano",
            "Tomado",
        ),
        (
            "YV-0125",
            dt.datetime(2026, 10, 9, 21, 0),
            None,
            "Casa Nogal",
            "transferencia",
            None,
            "Cancelado",
        ),
    ]
    assert wb["Pedidos"]["B2"].number_format == "dd/mm/yyyy hh:mm"

    docs = await outbox(sf)
    assert [(d.idempotency_key, d.message_type, d.template_key, d.recipient_jid) for d in docs] == [
        (f"daily_report:2026-10-09:v1:{p}", "document", "admin_daily_report", f"{p}@s.whatsapp.net")
        for p in ADMINS
    ]
    assert docs[0].body == (
        "📊 Reporte del 09/10/2026: 3 pedidos publicados, 2 tomados, 1 cancelados, "
        "2 moteros con actividad."
    )
    current = await report(sf)
    assert current is not None and (current.status, current.current_revision) == ("generating", 0)

    assert await deliver(ctx) == ["sent", "sent"]
    current = await report(sf)
    assert current is not None and (current.status, current.current_revision) == ("sent", 1)
    [rev] = await revisions(sf)
    assert rev.status == "sent"
    (media, _) = [json.loads(c.request.content) for c in evo["sendMedia"].calls]
    assert media["fileName"] == "reporte-ya-voy-2026-10-09.xlsx"

    # Replays do nothing: no new revision, no new message.
    assert await run(ctx) == Outcome.SKIPPED
    assert await report_jobs.process_daily_reports(ctx) == {}  # nothing left to look at
    assert len(await outbox(sf)) == 2


async def test_waits_for_pending_orders_and_reminds_once(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    sf = ctx["session_factory"]
    await world.order("YV-0056", "assigned", courier="Joel")
    pending = await world.order("YV-0057", "published", created=at(9, 22, 50))
    await world.order("YV-0058", "published", created=at(9, 22, 55))

    clock.now = at(9, 23)
    assert await run(ctx) == Outcome.WAITING
    current = await report(sf)
    assert current is not None and (current.status, current.pending_orders_count) == ("waiting", 2)
    clock.now = at(10, 0, 59)
    assert await run(ctx) == Outcome.WAITING
    assert await outbox(sf) == []

    clock.now = at(10, 1)  # due_at + REPORT_PENDING_REMINDER_MINUTES
    assert await run(ctx) == Outcome.WAITING
    reminders = await outbox(sf, "admin_report_pending_reminder")
    assert [r.idempotency_key for r in reminders] == [
        f"admin_report_pending_reminder:2026-10-09:20261010T020000Z:{p}" for p in ADMINS
    ]
    assert reminders[0].body == (
        "⏳ El reporte del 09/10/2026 está listo, pero lo retengo porque quedan pedidos sin "
        "tomar: YV-0057, YV-0058. Cuando estén tomados o cancelados te lo mando. Para cancelar "
        "uno: *cancelar YV-0057*."
    )
    clock.now = at(10, 3)
    assert await run(ctx) == Outcome.WAITING
    assert len(await outbox(sf)) == 2  # once

    await world.claim(pending, "Emiliano", at(10, 3, 10))
    async with world.engine.begin() as conn:
        await conn.execute(text("UPDATE orders SET status = 'cancelled' WHERE code = 'YV-0058'"))
    clock.now = at(10, 3, 30)
    assert await run(ctx) == Outcome.STARTED
    [rev] = await revisions(sf)
    assert [r[0] for r in rows(workbook(rev)["Pedidos"])[1:]] == ["YV-0056", "YV-0057", "YV-0058"]


async def test_only_drafts_pending_hold_the_report_without_a_reminder(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    await world.order(None, "awaiting_confirmation", created=at(9, 22, 58))
    clock.now = at(10, 2)
    assert await run(ctx) == Outcome.WAITING
    assert await outbox(ctx["session_factory"]) == []


async def test_empty_day_still_sends_the_file(ctx: dict[str, Any], clock: Clock) -> None:
    sf = ctx["session_factory"]
    clock.now = at(9, 23, 1)
    assert await run(ctx) == Outcome.STARTED
    [rev] = await revisions(sf)
    wb = workbook(rev)
    assert rows(wb["Resumen"]) == [("Motero", "Teléfono", "Tomados", "Liberados", "Cancelados")]
    assert len(rows(wb["Pedidos"])) == 1
    assert (await outbox(sf))[0].body == (
        "📊 Reporte del 09/10/2026: 0 pedidos publicados, 0 tomados, 0 cancelados, "
        "0 moteros con actividad."
    )


async def test_closed_day_lists_the_schedule_changes(
    ctx: dict[str, Any], clock: Clock, engine: AsyncEngine
) -> None:
    await f.insert_ok(
        engine,
        f.override(
            kind="closed_today",
            status="active",
            starts_at=at(9, 10),
            ends_at=at(10, 5),
            confirmed_at=at(9, 10),
            reason="lluvia",
        ),
    )
    await f.insert_ok(  # yesterday's closure does not belong to D
        engine,
        f.override(
            kind="closed_today",
            status="active",
            business_date=D - dt.timedelta(days=1),
            starts_at=at(8, 10),
            ends_at=at(9, 5),
            confirmed_at=at(8, 10),
        ),
    )
    clock.now = at(9, 23)
    assert await run(ctx) == Outcome.STARTED
    assert (await outbox(ctx["session_factory"]))[0].body == (
        "📊 Reporte del 09/10/2026: 0 pedidos publicados, 0 tomados, 0 cancelados, "
        "0 moteros con actividad.\nCambios de horario: cierre de hoy por lluvia"
    )


# --- due time from the 7a schedule core ----------------------------------------------------------


async def test_extension_past_midnight_moves_the_report(
    ctx: dict[str, Any], clock: Clock, world: World, engine: AsyncEngine
) -> None:
    sf = ctx["session_factory"]
    clock.now = at(9, 22, 30)
    assert await run(ctx) == Outcome.NOT_DUE
    await opening(engine, at(9, 23), at(10, 1), confirmed=at(9, 22, 40))
    await world.order(
        "YV-0130", "assigned", courier="Joel", created=at(10, 0, 30), assigned=at(10, 0, 40)
    )
    clock.now = at(9, 23, 30)
    assert await run(ctx) == Outcome.NOT_DUE
    clock.now = at(10, 1)
    assert await run(ctx) == Outcome.STARTED
    current = await report(sf)
    assert current is not None and current.due_at == at(10, 1)
    [rev] = await revisions(sf)
    [_, row] = rows(workbook(rev)["Pedidos"])
    assert (row[0], row[1], row[2]) == (
        "YV-0130",
        dt.datetime(2026, 10, 10, 0, 30),
        dt.datetime(2026, 10, 10, 0, 40),
    )
    assert (await outbox(sf))[0].body.endswith(
        "\nCambios de horario: apertura del 09/10/2026 de 23:00 a 01:00 (del 10/10)"
    )


async def test_waiting_report_follows_a_later_opening(
    ctx: dict[str, Any], clock: Clock, world: World, engine: AsyncEngine
) -> None:
    """Report waiting on a pending order at 23:00; an opening until 01:00 is confirmed at 23:10:
    ``due_at`` is recomputed while waiting (CU-16 1a), so the report waits for 01:00."""
    pending = await world.order("YV-0057", "published", created=at(9, 22, 55))
    clock.now = at(9, 23)
    assert await run(ctx) == Outcome.WAITING
    await opening(engine, at(9, 23), at(10, 1), confirmed=at(9, 23, 10))
    await world.claim(pending, "Joel", at(9, 23, 15))
    clock.now = at(9, 23, 20)
    assert await run(ctx) == Outcome.NOT_DUE
    current = await report(ctx["session_factory"])
    assert current is not None and (current.status, current.due_at) == ("waiting", at(10, 1))
    clock.now = at(10, 1)
    assert await run(ctx) == Outcome.STARTED


async def test_closure_confirmed_after_an_opening_does_not_move_the_due_time(
    ctx: dict[str, Any], clock: Clock, engine: AsyncEngine
) -> None:
    """7a precedence: the later closed_today wins over the opening (closed from 18:00), but
    ``report_due_at`` never moves earlier, so the report still waits for 01:00."""
    await opening(engine, at(9, 23), at(10, 1), confirmed=at(9, 10))
    await f.insert_ok(
        engine,
        f.override(
            kind="closed_today",
            status="active",
            starts_at=at(9, 18),
            ends_at=at(10, 5),
            confirmed_at=at(9, 18),
        ),
    )
    clock.now = at(9, 23, 30)
    assert await run(ctx) == Outcome.NOT_DUE
    clock.now = at(10, 1)
    assert await run(ctx) == Outcome.STARTED
    assert (await outbox(ctx["session_factory"]))[0].body.endswith(
        "\nCambios de horario: cierre de hoy, apertura del 09/10/2026 de 23:00 a 01:00 (del 10/10)"
    )


# --- business day rollover at 05:00 ---------------------------------------------------------------


async def test_orders_around_0500_and_catch_up_after_downtime(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    """The worker was down from 22:00 to 06:00. Stored UTC 07:59 / 08:01 of 10/10 are 04:59 /
    05:01 in Buenos Aires: the first order belongs to 09/10 (``business_date``), the second to
    10/10. At 06:00 the 09/10 report is produced; 10/10 is not due yet."""
    utc = dt.UTC
    await world.order(
        "YV-0201",
        "assigned",
        courier="Joel",
        created=dt.datetime(2026, 10, 10, 7, 59, tzinfo=utc),
        assigned=dt.datetime(2026, 10, 10, 7, 59, 30, tzinfo=utc),
    )
    await world.order(
        "YV-0202",
        "assigned",
        day=D1,
        courier="Emiliano",
        created=dt.datetime(2026, 10, 10, 8, 1, tzinfo=utc),
    )
    clock.now = at(10, 6)
    assert await report_jobs.process_daily_reports(ctx) == {
        "2026-10-09": "started",
        "2026-10-10": "not_due",
    }
    [rev] = await revisions(ctx["session_factory"])
    [_, row] = rows(workbook(rev)["Pedidos"])
    assert (row[0], row[1], row[2]) == (
        "YV-0201",
        dt.datetime(2026, 10, 10, 4, 59),
        dt.datetime(2026, 10, 10, 4, 59, 30),
    )
    assert await report(ctx["session_factory"], D1) is None


async def test_report_due_just_before_rollover(
    ctx: dict[str, Any], clock: Clock, engine: AsyncEngine
) -> None:
    """An opening until 02:00 (the latest close) and a check at 04:59 / 05:00."""
    await opening(engine, at(9, 23), at(10, 2), confirmed=at(9, 20))
    clock.now = at(10, 1, 59)
    assert await run(ctx) == Outcome.NOT_DUE
    clock.now = at(10, 4, 59)
    assert await report_jobs.process_daily_reports(ctx) == {"2026-10-09": "started"}
    clock.now = at(10, 5, 0)
    # still generating (document not delivered yet): the cron keeps watching it
    assert await report_jobs.process_daily_reports(ctx) == {
        "2026-10-09": "in_flight",
        "2026-10-10": "not_due",
    }
    assert set(await deliver(ctx)) == {"sent"}
    assert await report_jobs.process_daily_reports(ctx) == {"2026-10-10": "not_due"}


# --- CU-17: corrections ---------------------------------------------------------------------------


async def sent_report(ctx: dict[str, Any], clock: Clock) -> None:
    clock.now = at(9, 23)
    assert await run(ctx) == Outcome.STARTED
    assert set(await deliver(ctx)) == {"sent"}


async def test_release_then_other_courier_sends_a_correction(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    sf = ctx["session_factory"]
    order = await world.order("YV-0123", "assigned", courier="Joel", created=at(9, 21))
    await world.order("YV-0124", "assigned", courier="Joel", created=at(9, 21, 30))
    await sent_report(ctx, clock)

    async with sf() as session, session.begin():
        await world.release(session, order, "Joel")
    current = await report(sf)
    assert current is not None and current.status == "correction_pending"

    clock.now = at(9, 23, 40)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "waiting"}
    clock.now = at(10, 2)  # no reminder for corrections caused by a release (CU-17)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "waiting"}
    assert await outbox(sf, "admin_report_pending_reminder") == []

    await world.claim(order, "Emiliano", at(10, 2, 5))
    clock.now = at(10, 2, 10)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "started"}
    [_, rev2] = await revisions(sf)
    assert (rev2.revision, rev2.kind, rev2.changes) == (
        2,
        "correction",
        [{"code": "YV-0123", "type": "courier_changed", "from": "Joel", "to": "Emiliano"}],
    )
    assert rev2.file_name == "reporte-ya-voy-2026-10-09-correccion-2.xlsx"
    docs = await outbox(sf, "admin_daily_report_correction")
    assert [d.idempotency_key for d in docs] == [f"daily_report:2026-10-09:v2:{p}" for p in ADMINS]
    assert docs[0].body == (
        "🛠️ *Corrección* del reporte del 09/10/2026 (versión 2):\n"
        "• El pedido *YV-0123* pasó de Joel a Emiliano.\n\n"
        "Totales actualizados: 2 pedidos publicados, 2 tomados, 0 cancelados. "
        "Este archivo reemplaza a la versión anterior."
    )
    wb = workbook(rev2)
    assert wb.sheetnames == ["Resumen", "Pedidos", "Correcciones"]
    assert rows(wb["Correcciones"]) == [
        ("CORRECCIÓN — versión 2", None, None, None),
        ("Código", "Antes", "Ahora", "Detalle"),
        (
            "YV-0123",
            "Tomado por Joel",
            "Tomado por Emiliano",
            "El pedido YV-0123 pasó de Joel a Emiliano.",
        ),
    ]
    pedidos = rows(wb["Pedidos"])
    assert pedidos[0][0] == "CORRECCIÓN — versión 2"
    assert [(r[0], r[2], r[8], r[10]) for r in pedidos[2:]] == [
        ("YV-0123", dt.datetime(2026, 10, 10, 2, 5), "Emiliano", "Cambió de motero"),
        ("YV-0124", dt.datetime(2026, 10, 9, 21, 33), "Joel", None),
    ]
    assert rows(wb["Resumen"])[2:] == [
        ("Emiliano", None, 1, 0, 0),
        ("Joel", "5491100000007", 1, 1, 0),
    ]

    assert set(await deliver(ctx)) == {"sent"}
    current = await report(sf)
    assert current is not None and (current.status, current.current_revision) == ("sent", 2)


async def test_release_and_retake_by_the_same_courier_is_no_correction(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    sf = ctx["session_factory"]
    order = await world.order("YV-0123", "assigned", courier="Joel")
    await sent_report(ctx, clock)
    async with sf() as session, session.begin():
        await world.release(session, order, "Joel")
    await world.claim(order, "Joel", at(9, 23, 45))
    clock.now = at(9, 23, 50)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "no_changes"}
    current = await report(sf)
    assert current is not None
    assert (current.status, current.current_revision, current.correction_requested_at) == (
        "sent",
        1,
        None,
    )
    assert len(await revisions(sf)) == 1
    assert len(await outbox(sf)) == 2


async def test_cancel_after_report(ctx: dict[str, Any], clock: Clock, world: World) -> None:
    sf = ctx["session_factory"]
    order = await world.order("YV-0123", "assigned", courier="Joel")
    await sent_report(ctx, clock)
    async with sf() as session, session.begin():
        await world.cancel(session, order, "Joel")
    clock.now = at(9, 23, 31)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "started"}
    [doc, _] = await outbox(sf, "admin_daily_report_correction")
    assert "• El pedido *YV-0123* fue cancelado (estaba asignado a Joel).\n" in doc.body
    [_, rev2] = await revisions(sf)
    assert rows(workbook(rev2)["Resumen"])[2:] == [("Joel", "5491100000007", 0, 0, 1)]


async def test_opening_after_report_adds_orders_with_reminder(
    ctx: dict[str, Any], clock: Clock, world: World, engine: AsyncEngine
) -> None:
    """RF-55: reopened by an opening (docs/03 SQL, as 7a does on confirmation); the correction
    waits for the new due time and gets the reminder variant."""
    sf = ctx["session_factory"]
    await world.order("YV-0123", "assigned", courier="Joel")
    await sent_report(ctx, clock)
    override_id = await opening(engine, at(9, 23), at(10, 1), confirmed=at(9, 23, 30))
    async with engine.begin() as conn:
        await conn.execute(
            text(
                "UPDATE daily_reports SET status = 'correction_pending', due_at = :due, "
                "reopened_by_override_id = :id, pending_reminder_sent_at = NULL, "
                "correction_requested_at = COALESCE(correction_requested_at, now()) "
                "WHERE report_date = :d AND status IN ('sent', 'generating', 'correction_pending')"
            ),
            {"due": at(10, 1), "id": override_id, "d": D},
        )
    o130 = await world.order("YV-0130", "assigned", courier="Joel", created=at(9, 23, 40))
    await world.order("YV-0131", "cancelled", created=at(9, 23, 50))
    late = await world.order("YV-0132", "published", created=at(10, 0, 55))

    clock.now = at(10, 0, 59)
    assert await report_jobs.process_report_corrections(ctx) == {}  # not due yet
    clock.now = at(10, 1)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "waiting"}
    clock.now = at(10, 3)
    await report_jobs.process_report_corrections(ctx)
    [reminder, _] = await outbox(sf, "admin_report_pending_reminder")
    assert reminder.body == (
        "⏳ La corrección del reporte del 09/10/2026 está lista, pero la retengo porque quedan "
        "pedidos sin tomar: YV-0132. Cuando estén tomados o cancelados te la mando. Para "
        "cancelar uno: *cancelar YV-0132*."
    )
    await world.claim(late, "Emiliano", at(10, 3, 5))
    clock.now = at(10, 3, 10)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "started"}
    [doc, _] = await outbox(sf, "admin_daily_report_correction")
    assert doc.body.split("\n\n")[0] == (
        "🛠️ *Corrección* del reporte del 09/10/2026 (versión 2):\n"
        "• Pedido agregado *YV-0130* (tomado por Joel)\n"
        "• Pedido agregado *YV-0131* (cancelado)\n"
        "• Pedido agregado *YV-0132* (tomado por Emiliano)"
    )
    assert o130
    await deliver(ctx)
    current = await report(sf)
    assert current is not None
    assert (current.status, current.current_revision, current.reopened_by_override_id) == (
        "sent",
        2,
        None,
    )


async def test_opening_after_report_without_new_orders_sends_nothing(
    ctx: dict[str, Any], clock: Clock, world: World, engine: AsyncEngine
) -> None:
    sf = ctx["session_factory"]
    await world.order("YV-0123", "assigned", courier="Joel")
    await sent_report(ctx, clock)
    override_id = await opening(engine, at(9, 23), at(10, 1), confirmed=at(9, 23, 30))
    async with engine.begin() as conn:
        await conn.execute(
            update(m.DailyReport).values(
                status="correction_pending", due_at=at(10, 1), reopened_by_override_id=override_id
            )
        )
    clock.now = at(10, 1)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "no_changes"}
    current = await report(sf)
    assert current is not None
    assert (current.status, current.reopened_by_override_id) == ("sent", None)
    assert len(await outbox(sf)) == 2


# --- idempotency and races -----------------------------------------------------------------------


async def test_concurrent_cron_runs_produce_one_report(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    await world.order("YV-0123", "assigned", courier="Joel")
    clock.now = at(9, 23)
    results = await asyncio.gather(*(report_jobs.process_daily_reports(ctx) for _ in range(6)))
    outcomes = sorted(r["2026-10-09"] for r in results)
    assert outcomes.count("started") == 1, outcomes
    assert len(await revisions(ctx["session_factory"])) == 1
    assert len(await outbox(ctx["session_factory"])) == 2


async def test_concurrent_correction_runs_produce_one_revision(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    sf = ctx["session_factory"]
    order = await world.order("YV-0123", "assigned", courier="Joel")
    await sent_report(ctx, clock)
    async with sf() as session, session.begin():
        await world.cancel(session, order, "Joel")
    clock.now = at(9, 23, 40)
    results = await asyncio.gather(
        *(report_jobs.process_report_corrections(ctx) for _ in range(3)),
        *(report_jobs.process_daily_reports(ctx) for _ in range(3)),
    )
    assert [r.get("2026-10-09") for r in results].count("started") == 1
    assert [r.revision for r in await revisions(sf)] == [1, 2]
    assert len(await outbox(sf, "admin_daily_report_correction")) == 2


async def test_report_row_held_by_another_run_is_skipped(
    ctx: dict[str, Any], clock: Clock, engine: AsyncEngine
) -> None:
    clock.now = at(9, 23)
    await f.insert_ok(engine, f.daily_report(due_at=at(9, 23)))
    async with engine.connect() as holder:
        await holder.begin()
        await holder.execute(text("SELECT 1 FROM daily_reports FOR UPDATE"))
        assert await run(ctx) == Outcome.BUSY
        await holder.rollback()
    assert await run(ctx) == Outcome.STARTED


async def test_release_in_flight_makes_the_report_wait(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    """A release not yet committed holds the order row: the report does not read a state about
    to change (NOWAIT, retried on the next run), then sees the order pending."""
    sf = ctx["session_factory"]
    order = await world.order("YV-0123", "assigned", courier="Joel")
    clock.now = at(9, 23)
    async with sf() as releasing:
        await releasing.begin()
        await world.release(releasing, order, "Joel")
        assert await run(ctx) == Outcome.BUSY
        await releasing.commit()
    assert await run(ctx) == Outcome.WAITING
    assert await revisions(sf) == []


async def wait_for_lock_wait(engine: AsyncEngine) -> None:
    for _ in range(200):
        async with engine.connect() as conn:
            waiting = await conn.scalar(
                text("SELECT count(*) FROM pg_stat_activity WHERE wait_event_type = 'Lock'")
            )
        if waiting:
            return
        await asyncio.sleep(0.02)
    raise AssertionError("no transaction is waiting for a lock")


async def test_release_waiting_on_the_snapshot_triggers_a_correction(
    ctx: dict[str, Any], clock: Clock, world: World, engine: AsyncEngine
) -> None:
    """The cron took the snapshot (orders share-locked, report generating) before the release:
    the release waits for the commit, then finds the report ``generating`` and asks for a
    correction, which goes out once the original is delivered and the order resolved."""
    sf = ctx["session_factory"]
    settings = ctx["settings"]
    order = await world.order("YV-0123", "assigned", courier="Joel")
    clock.now = at(9, 23)
    async with sf() as cron:
        await cron.begin()
        result = await service.step(cron, settings, D)
        assert result.outcome == Outcome.STARTED and result.job is not None

        async def release() -> None:
            async with sf() as session, session.begin():
                await world.release(session, order, "Joel")

        releasing = asyncio.create_task(release())
        await wait_for_lock_wait(engine)
        assert not releasing.done()
        await cron.commit()
    await releasing
    current = await report(sf)
    assert current is not None and current.status == "correction_pending"

    async with sf() as session, session.begin():
        ids = await service.store_file(session, settings, result.job, b"PK-file")
    assert len(ids) == 2
    assert set(await deliver(ctx)) == {"sent"}
    current = await report(sf)
    assert current is not None and (current.status, current.current_revision) == (
        "correction_pending",
        1,
    )

    clock.now = at(9, 23, 5)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "waiting"}
    await world.claim(order, "Emiliano", at(9, 23, 6))
    clock.now = at(9, 23, 7)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "started"}
    [_, rev2] = await revisions(sf)
    assert rev2.changes == [
        {"code": "YV-0123", "type": "courier_changed", "from": "Joel", "to": "Emiliano"}
    ]


async def test_change_while_a_correction_is_in_flight_is_not_lost(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    sf = ctx["session_factory"]
    first = await world.order("YV-0123", "assigned", courier="Joel")
    second = await world.order("YV-0124", "assigned", courier="Joel")
    await sent_report(ctx, clock)
    async with sf() as session, session.begin():
        await world.cancel(session, first, "Joel")
    clock.now = at(9, 23, 40)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "started"}
    async with sf() as session, session.begin():  # v2 generated, not yet delivered
        await world.cancel(session, second, "Joel")
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "in_flight"}
    assert set(await deliver(ctx)) == {"sent"}
    current = await report(sf)
    assert current is not None and (current.status, current.current_revision) == (
        "correction_pending",
        2,
    )
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "started"}
    [_, _, rev3] = await revisions(sf)
    assert rev3.changes == [{"code": "YV-0124", "type": "cancelled", "from": "Joel", "to": None}]


# --- failures and retries ---------------------------------------------------------------------


async def test_build_failure_is_retried_then_reported(ctx: dict[str, Any], clock: Clock) -> None:
    sf = ctx["session_factory"]
    settings = ctx["settings"]

    def broken(data: Any, tz: Any) -> bytes:
        raise ValueError("boom")

    clock.now = at(9, 23)
    for attempt in (1, 2):
        assert await service.run_step(sf, settings, D, builder=broken) == (Outcome.STARTED, [])
        [rev] = await revisions(sf)
        assert (rev.status, rev.attempts, rev.last_error) == ("failed", attempt, "ValueError: boom")
    _, ids = await service.run_step(sf, settings, D, builder=broken)
    assert len(ids) == 2
    current = await report(sf)
    assert current is not None and current.status == "failed"
    [alert, _] = await outbox(sf, "admin_daily_report_failed")
    assert alert.idempotency_key == f"admin_daily_report_failed:2026-10-09:v1:{ADMINS[0]}"
    assert alert.body == (
        "❗ No pude generar o enviar el reporte del 09/10/2026 (versión 1). "
        "Error: no pude armar el archivo."
    )
    assert await run(ctx) == Outcome.SKIPPED  # terminal


async def test_build_failure_then_success(ctx: dict[str, Any], clock: Clock) -> None:
    sf = ctx["session_factory"]

    def broken(data: Any, tz: Any) -> bytes:
        raise RuntimeError("disk full")

    clock.now = at(9, 23)
    await service.run_step(sf, ctx["settings"], D, builder=broken)
    assert await run(ctx) == Outcome.STARTED
    [rev] = await revisions(sf)
    assert (rev.status, rev.attempts, rev.last_error) == ("generated", 2, None)


async def test_stale_generation_is_retried_and_the_old_attempt_discarded(
    ctx: dict[str, Any], clock: Clock
) -> None:
    sf = ctx["session_factory"]
    settings = ctx["settings"]
    clock.now = at(9, 23)
    async with sf() as session, session.begin():
        stalled = await service.step(session, settings, D)  # the worker dies before building
    assert stalled.job is not None
    clock.now = at(9, 23, 9)
    assert await run(ctx) == Outcome.IN_FLIGHT
    clock.now = at(9, 23, 11)
    assert await run(ctx) == Outcome.STARTED
    [rev] = await revisions(sf)
    assert (rev.status, rev.attempts) == ("generated", 2)
    async with sf() as session, session.begin():
        assert await service.store_file(session, settings, stalled.job, b"late") == []
    assert len(await outbox(sf)) == 2


async def test_delivery_failure_fails_the_report_with_one_alert(
    ctx: dict[str, Any], clock: Clock, evo: respx.MockRouter
) -> None:
    sf = ctx["session_factory"]
    clock.now = at(9, 23)
    assert await run(ctx) == Outcome.STARTED
    evo["sendMedia"].side_effect = None
    evo["sendMedia"].respond(400, json={"error": "bad"})
    assert await deliver(ctx) == ["failed", "failed"]
    current = await report(sf)
    assert current is not None and current.status == "failed"
    [rev] = await revisions(sf)
    assert rev.status == "failed"
    alerts = await outbox(sf, "admin_daily_report_failed")
    assert len(alerts) == 2  # one per admin, not one per failed copy
    assert alerts[0].body == (
        "❗ No pude generar o enviar el reporte del 09/10/2026 (versión 1). "
        "Error: no se pudo enviar el archivo a un administrador tras 1 intento."
    )
    assert await outbox(sf, "admin_outbound_failed") == []


async def test_mark_report_for_correction_only_after_a_send(
    session_factory: SF, engine: AsyncEngine
) -> None:
    for status, expected in (
        ("waiting", "waiting"),
        ("generating", "correction_pending"),
        ("sent", "correction_pending"),
        ("failed", "failed"),
    ):
        async with engine.begin() as conn:
            await conn.execute(text("TRUNCATE daily_reports CASCADE"))
        await f.insert_ok(engine, f.daily_report(status=status))
        async with session_factory() as session, session.begin():
            await service.mark_report_for_correction(session, D)
        current = await report(session_factory)
        assert current is not None and current.status == expected, status


# --- manual resend of failed reports (make resend-report) --------------------------------------

D0 = D - dt.timedelta(days=1)


def report_errors(caplog: pytest.LogCaptureFixture) -> list[logging.LogRecord]:
    return [
        r
        for r in caplog.records
        if r.name == "ya_voy.reports.service" and r.levelno == logging.ERROR
    ]


async def failed_delivery(ctx: dict[str, Any], clock: Clock, evo: respx.MockRouter) -> None:
    """Revision 1 built; WhatsApp rejects both copies; the report is ``failed``."""
    clock.now = at(9, 23)
    assert await run(ctx) == Outcome.STARTED
    evo["sendMedia"].side_effect = None
    evo["sendMedia"].respond(400, json={"error": "disconnected"})
    await deliver(ctx)
    current = await report(ctx["session_factory"])
    assert current is not None and current.status == "failed"


def evolution_back(evo: respx.MockRouter) -> None:
    evo["sendMedia"].respond(201, json={"key": {"id": "BAE5RESENT", "fromMe": True}})


async def document_rows(sf: SF) -> list[m.OutboundMessage]:
    return [r for r in await outbox(sf) if r.message_type == "document"]


async def test_delivery_failure_is_logged_at_error_per_recipient(
    ctx: dict[str, Any], clock: Clock, evo: respx.MockRouter, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level(logging.INFO):
        await failed_delivery(ctx, clock, evo)
    errors = report_errors(caplog)
    assert len(errors) == 2  # one per failed copy
    for record in errors:
        message = record.getMessage()
        assert message.startswith(
            "daily report 2026-10-09 v1 FAILED: delivery to un administrador failed after 1 "
        )
        assert "disconnected" in message
        assert (record.report_date, record.revision, record.recipient) == (  # type: ignore[attr-defined]
            "2026-10-09",
            1,
            "un administrador",
        )
        assert "disconnected" in record.error  # type: ignore[attr-defined]
    assert {r.outbound_message_id for r in errors} == {  # type: ignore[attr-defined]
        r.id for r in await document_rows(ctx["session_factory"])
    }


async def test_resend_after_delivery_failure_requeues_the_same_file(
    ctx: dict[str, Any], clock: Clock, evo: respx.MockRouter
) -> None:
    sf = ctx["session_factory"]
    settings = ctx["settings"]
    await failed_delivery(ctx, clock, evo)
    [rev] = await revisions(sf)
    original = rev.file_content
    evolution_back(evo)
    clock.now = at(10, 8)

    outcomes, ids = await service.resend_failed_reports(sf, settings, [D])

    assert outcomes == {"2026-10-09": "resent"}
    assert len(ids) == 2
    [rev] = await revisions(sf)
    assert (rev.status, rev.resends, rev.attempts, rev.file_content) == (
        "generated",
        1,
        1,
        original,
    )
    current = await report(sf)
    assert current is not None and current.status == "generating"
    resent = [r for r in await document_rows(sf) if r.id in ids]
    assert sorted(r.idempotency_key for r in resent) == [
        f"daily_report:2026-10-09:v1:r1:{phone}" for phone in ADMINS
    ]
    assert {r.template_key for r in resent} == {"admin_daily_report"}

    # Running it again does nothing: the report is no longer failed.
    assert await service.resend_failed_reports(sf, settings, [D]) == (
        {"2026-10-09": "not_failed"},
        [],
    )

    await deliver(ctx)
    [rev] = await revisions(sf)
    assert rev.status == "sent"
    current = await report(sf)
    assert current is not None
    # The file may predate changes made while it was failed: one more comparison.
    assert (current.status, current.current_revision) == ("correction_pending", 1)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "no_changes"}
    current = await report(sf)
    assert current is not None and current.status == "sent"
    assert len(await revisions(sf)) == 1


async def test_resend_skips_admins_that_already_got_the_file(
    ctx: dict[str, Any], clock: Clock, evo: respx.MockRouter
) -> None:
    sf = ctx["session_factory"]
    clock.now = at(9, 23)
    assert await run(ctx) == Outcome.STARTED
    evo["sendMedia"].side_effect = [
        httpx.Response(201, json={"key": {"id": "BAE5OK", "fromMe": True}}),
        httpx.Response(400, json={"error": "disconnected"}),
    ]
    await deliver(ctx)
    current = await report(sf)
    assert current is not None and current.status == "failed"
    evo["sendMedia"].side_effect = None
    evolution_back(evo)

    outcomes, ids = await service.resend_failed_reports(sf, ctx["settings"], [D])

    assert outcomes == {"2026-10-09": "resent"}
    [row] = [r for r in await document_rows(sf) if r.id in ids]
    assert row.idempotency_key == f"daily_report:2026-10-09:v1:r1:{ADMINS[1]}"
    await deliver(ctx)
    [rev] = await revisions(sf)
    assert rev.status == "sent"


async def test_failed_resend_fails_again_with_a_new_alert(
    ctx: dict[str, Any], clock: Clock, evo: respx.MockRouter, caplog: pytest.LogCaptureFixture
) -> None:
    sf = ctx["session_factory"]
    await failed_delivery(ctx, clock, evo)
    await deliver(ctx)  # the first alerts
    _, ids = await service.resend_failed_reports(sf, ctx["settings"], [D])
    assert len(ids) == 2
    caplog.clear()
    with caplog.at_level(logging.INFO):
        await deliver(ctx)  # WhatsApp still down
    assert len(report_errors(caplog)) == 2
    current = await report(sf)
    assert current is not None and current.status == "failed"
    [rev] = await revisions(sf)
    assert (rev.status, rev.resends) == ("failed", 1)
    keys = sorted(r.idempotency_key for r in await outbox(sf, "admin_daily_report_failed"))
    assert keys == sorted(
        [f"admin_daily_report_failed:2026-10-09:v1:{p}" for p in ADMINS]
        + [f"admin_daily_report_failed:2026-10-09:v1:r1:{p}" for p in ADMINS]
    )
    # A second resend round gets its own keys.
    evolution_back(evo)
    _, ids = await service.resend_failed_reports(sf, ctx["settings"], [D])
    rows = [r for r in await document_rows(sf) if r.id in ids]
    assert sorted(r.idempotency_key for r in rows) == [
        f"daily_report:2026-10-09:v1:r2:{p}" for p in ADMINS
    ]


async def broken_build(ctx: dict[str, Any], day: dt.date = D) -> None:
    def broken(data: Any, tz: Any) -> bytes:
        raise OSError("disk full")

    for _ in range(ctx["settings"].daily_report_max_attempts):
        await service.run_step(ctx["session_factory"], ctx["settings"], day, builder=broken)
    current = await report(ctx["session_factory"], day)
    assert current is not None and current.status == "failed"


async def test_build_failures_log_warnings_then_one_error(
    ctx: dict[str, Any], clock: Clock, caplog: pytest.LogCaptureFixture
) -> None:
    clock.now = at(9, 23)
    with caplog.at_level(logging.INFO):
        await broken_build(ctx)
    warnings = [
        r
        for r in caplog.records
        if r.name == "ya_voy.reports.service" and r.levelno == logging.WARNING
    ]
    assert [r.getMessage() for r in warnings] == [
        "daily report 2026-10-09 v1: build attempt 1 failed: OSError: disk full",
        "daily report 2026-10-09 v1: build attempt 2 failed: OSError: disk full",
    ]
    [error] = report_errors(caplog)
    assert error.getMessage() == (
        "daily report 2026-10-09 v1 FAILED: build failed after 3 attempts: OSError: disk full"
    )
    assert (error.report_date, error.revision, error.error) == (  # type: ignore[attr-defined]
        "2026-10-09",
        1,
        "OSError: disk full",
    )


async def test_resend_rebuilds_a_report_without_file(ctx: dict[str, Any], clock: Clock) -> None:
    sf = ctx["session_factory"]
    clock.now = at(9, 23)
    await broken_build(ctx)
    clock.now = at(10, 8)

    outcomes, ids = await service.resend_failed_reports(sf, ctx["settings"], [D])

    assert outcomes == {"2026-10-09": "rebuild:started"}
    [rev] = await revisions(sf)
    assert (rev.status, rev.attempts, rev.resends, rev.last_error) == ("generated", 1, 1, None)
    assert rev.file_content is not None
    rows = [r for r in await document_rows(sf) if r.id in ids]
    assert sorted(r.idempotency_key for r in rows) == [
        f"daily_report:2026-10-09:v1:r1:{p}" for p in ADMINS
    ]
    await deliver(ctx)
    assert await report_jobs.process_report_corrections(ctx) == {"2026-10-09": "no_changes"}
    current = await report(sf)
    assert current is not None and (current.status, current.current_revision) == ("sent", 1)


async def test_resend_rebuilds_a_failed_correction_with_current_data(
    ctx: dict[str, Any], clock: Clock, world: World
) -> None:
    sf = ctx["session_factory"]
    order_id = await world.order("YV-0123", "assigned", courier="Joel")
    await sent_report(ctx, clock)
    async with sf() as session, session.begin():
        await world.release(session, order_id, "Joel")
    await world.claim(order_id, "Emiliano", at(9, 23, 5))
    clock.now = at(9, 23, 6)
    await broken_build(ctx)
    [_, rev2] = await revisions(sf)
    assert (rev2.status, rev2.attempts) == ("failed", 3)

    outcomes, ids = await service.resend_failed_reports(sf, ctx["settings"], [D])

    assert outcomes == {"2026-10-09": "rebuild:started"}
    assert len(ids) == 2
    [_, rev2] = await revisions(sf)
    assert (rev2.revision, rev2.status, rev2.attempts, rev2.resends) == (2, "generated", 1, 1)
    assert rev2.changes == [
        {"code": "YV-0123", "type": "courier_changed", "from": "Joel", "to": "Emiliano"}
    ]


async def test_resend_all_failed_and_nothing_to_do(
    ctx: dict[str, Any], clock: Clock, evo: respx.MockRouter, engine: AsyncEngine
) -> None:
    sf = ctx["session_factory"]
    settings = ctx["settings"]
    clock.now = at(9, 23)
    await broken_build(ctx, D0)  # yesterday: build failed
    await failed_delivery(ctx, clock, evo)  # today: delivery failed
    await f.insert_ok(engine, f.daily_report(report_date=D1, status="sent"))  # untouched
    evolution_back(evo)

    outcomes, ids = await service.resend_failed_reports(sf, settings)

    assert outcomes == {"2026-10-08": "rebuild:started", "2026-10-09": "resent"}
    assert len(ids) == 4
    sent_row = await report(sf, D1)
    assert sent_row is not None and sent_row.status == "sent"
    # Nothing left: a second run (single date, unknown date or all) does nothing.
    assert await service.resend_failed_reports(sf, settings) == ({}, [])
    assert await service.resend_failed_reports(sf, settings, [D, D1, D + dt.timedelta(9)]) == (
        {"2026-10-09": "not_failed", "2026-10-10": "not_failed", "2026-10-18": "not_found"},
        [],
    )


async def test_concurrent_resends_queue_one_round(
    ctx: dict[str, Any], clock: Clock, evo: respx.MockRouter
) -> None:
    sf = ctx["session_factory"]
    await failed_delivery(ctx, clock, evo)
    results = await asyncio.gather(
        *(service.resend_failed_reports(sf, ctx["settings"], [D]) for _ in range(4))
    )
    assert sorted(o["2026-10-09"] for o, _ in results) == ["not_failed"] * 3 + ["resent"]
    assert sum(len(ids) for _, ids in results) == 2
    [rev] = await revisions(sf)
    assert rev.resends == 1
    assert len(await document_rows(sf)) == 4  # 2 originals + 1 round


async def test_concurrent_rebuild_resends_build_once(ctx: dict[str, Any], clock: Clock) -> None:
    sf = ctx["session_factory"]
    clock.now = at(9, 23)
    await broken_build(ctx)
    results = await asyncio.gather(
        *(service.resend_failed_reports(sf, ctx["settings"], [D]) for _ in range(4))
    )
    assert sum(len(ids) for _, ids in results) == 2
    [rev] = await revisions(sf)
    assert (rev.status, rev.attempts, rev.resends) == ("generated", 1, 1)
    assert len(await document_rows(sf)) == 2


async def test_resend_script_queues_the_documents(
    ctx: dict[str, Any], clock: Clock, evo: respx.MockRouter, database_url: str
) -> None:
    await failed_delivery(ctx, clock, evo)
    settings = ctx["settings"].model_copy(
        update={"database_url": database_url, "redis_url": redis_test_url()}
    )

    outcomes, ids = await resend_report.run(settings, [D], all_failed=False)
    assert outcomes == {"2026-10-09": "resent"}
    assert len(ids) == 2
    jobs = await queued(ctx["redis"], SEND_OUTBOUND)
    assert {args[0] for _, args in jobs} >= set(ids)

    assert await resend_report.run(settings, None, all_failed=True) == ({}, [])


def test_resend_script_needs_dates_or_all_failed() -> None:
    for argv in ([], ["2026-10-09", "--all-failed"], ["09/10/2026"]):
        with pytest.raises(SystemExit):
            resend_report.main(argv)
